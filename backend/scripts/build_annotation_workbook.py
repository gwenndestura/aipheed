"""
build_annotation_workbook.py
----------------------------
Generate the expert-annotation workbook for the audited CALABARZON
food-insecurity news dataset.

Supersedes build_review_sheet.py, which was written for the pre-audit
1,168-row file and the retired DOST Region IV-A framing. The design here
follows docs/annotator_selection_guide.md:

  * one workbook per reviewer, so reviewers never see each other's calls;
  * a BLIND set that hides the machine's labels, so agreement measured on
    that subset is not inflated by anchoring on the proposed tag;
  * the blind set silently mixes in articles the audit REMOVED, so the
    reviewers measure recall (wrongly-dropped articles) and not only
    precision (wrongly-kept ones);
  * an annotator-credentials sheet, so the thesis's validity annex is
    filled in at source rather than reconstructed afterwards;
  * identical questions on both sheets, so blind and anchored answers are
    directly comparable.

The blind set's true audit decision is written to a separate key file, not
into any reviewer's workbook.

Usage
-----
    python scripts/build_annotation_workbook.py                 # reviewers A, B
    python scripts/build_annotation_workbook.py -r A B C --pilot 12
"""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "processed" / "calabarzon_food_insecurity_dataset.parquet"
DROPS = ROOT / "data" / "processed" / "calabarzon_dataset_dropped_audit.csv"
STAGE1 = ROOT / "data" / "processed" / "_audit_stage1.parquet"
OUTDIR = ROOT / "data" / "processed" / "annotation"

# ---------------------------------------------------------------------------
# Label space — mirrors scripts/audit_stage2_rebuild.py, which is the source
# of truth. validate_label_space() fails the build if the dataset drifts past
# what this workbook can offer.
# ---------------------------------------------------------------------------
HYPOTHESIS_TEXT = {
    "T1": "Food supply disruption or food price increases affecting access to food",
    "T1b": "Fish kill or aquaculture collapse reducing fish food supply",
    "T2": "Health services or nutrition programs unavailable or unaffordable",
    "T3": "Government food security programs ineffective or unavailable",
    "T4": "Economic hardship reducing household income and food purchasing power",
    "T5": "Infrastructure failures limiting food transport or storage",
    "T6": "Agricultural land loss or conversion reducing food production",
    "T7": "Civil displacement or evacuation reducing food access",
    "T8": "Social unrest or conflict disrupting food systems",
    "T9": "OFW remittance reduction reducing household food purchasing power",
}
HYPOTHESIS_DIM = {
    "T1": "B", "T1b": "A", "T2": "C", "T3": "E", "T4": "B",
    "T5": "D", "T6": "A", "T7": "D", "T8": "D", "T9": "F",
}
HYPOTHESIS_GLOSS = {
    "T1": "food supply / price disruption",
    "T1b": "fish kill / aquaculture collapse",
    "T2": "health / nutrition services gap",
    "T3": "food programme ineffective or absent",
    "T4": "economic hardship / poverty",
    "T5": "transport / storage failure",
    "T6": "farmland loss / production drop",
    "T7": "displacement / evacuation",
    "T8": "unrest / conflict",
    "T9": "OFW remittance drop",
}
DIMENSIONS = {
    "A": ("Availability (production)",
          "Did something reduce the actual supply of food — growing it, catching it, "
          "raising it?",
          "Food at the source: farms, fishing grounds, livestock. Typhoon flattens rice "
          "fields, fish kill in Taal Lake, ASF wipes out hogs, pests, a failed harvest. "
          "Less food EXISTS because production was hit."),
    "B": ("Access / affordability",
          "Is the food there, but people cannot afford it or reach it?",
          "The wallet and the market. Rice prices jump, vegetables mahal, poverty or lost "
          "purchasing power limits what a household can buy. The food is available but "
          "out of reach."),
    "C": ("Utilization / nutrition",
          "Is it about health, nutrition, or malnutrition?",
          "What food does once eaten. Child malnutrition or stunting, feeding programmes, "
          "nutrition advisories, ENNS or FIES findings. Having food is not enough — is it "
          "keeping people healthy?"),
    "D": ("Stability (shocks)",
          "Did a disaster or disruption suddenly cut off food access?",
          "The interruption itself. A flood displaces families, roads are cut, a transport "
          "strike halts deliveries, unrest closes markets. A sudden event broke the normal "
          "flow of food."),
    "E": ("Hunger / assistance",
          "Is it about actual hunger, or about giving out food help?",
          "The outcome, or the response to it. Families going hungry, DSWD food packs and "
          "ayuda, relief feeding, a food programme that failed to reach the people it was "
          "meant for."),
    "F": ("Livelihood",
          "Did lost income or jobs threaten families' ability to buy food?",
          "Earning power specifically: jobs, wages, OFW remittances. Remittances drop, mass "
          "layoffs, fisherfolk lose their income source."),
}
SEVERITY = [
    ("High",
     "Widespread or acute. Many households, several barangays or LGUs, a large quantified "
     "loss, a calamity declaration with a stated food impact, or hunger reported as "
     "already happening."),
    ("Moderate",
     "Real but contained. One barangay or municipality, a moderate or partial loss, an "
     "ongoing strain the article does not describe as acute."),
    ("Low",
     "Minor, early, or precautionary. A small or anticipated impact, a warning or forecast, "
     "or a routine programme with no crisis reported."),
    ("Cannot tell from article",
     "The article does not give enough to judge scale. Use this freely — a blank is worth "
     "more to us than a guess."),
]
BOUNDARIES = [
    ("A vs B",
     "the loss of supply, or the price of what is left. \"Typhoon destroyed the rice crop\" "
     "= A. \"Rice prices rose after the typhoon\" = B."),
    ("A vs D",
     "the production loss, or the disruption. Crops or harvest destroyed = A. People "
     "evacuated, roads cut, deliveries stopped = D."),
    ("B vs F",
     "no money to buy, or no income earned. \"Food too expensive for families\" = B. "
     "\"Families lost jobs, remittances fell\" = F. In this scheme general economic "
     "hardship (T4) sits in B; F is for income and remittances specifically (T9)."),
    ("D vs E",
     "the displacement, or the response to it. Families evacuated and cut off = D. Food "
     "packs distributed to them = E. Ask which one the article is actually reporting."),
    ("E vs B",
     "hunger and aid, or affordability. Relief distribution and reported hunger = E, even "
     "though both are about getting food. B is specifically about price and purchasing "
     "power."),
]
LOCATION_TRAPS = [
    "Surnames read as towns — Tiu Laurel (Laurel, Batangas), Lopez (Lopez, Quezon), "
    "Pangilinan (Pangil, Laguna), Rodriguez (Rodriguez, Rizal).",
    "Foreign places — Laguna Beach and Laguna Hills in California; Nigeria's NAIC read "
    "as Naic, Cavite; Indonesian rice stories read as Rizal.",
    "Dateline instead of subject — the Inquirer's Lucena bureau datelines regional "
    "stories \"LUCENA CITY\" even when the subject is Rizal or Batangas.",
    "Same name, other region — Talisay City (Cebu and Batangas), San Juan (NCR and "
    "Batangas), Quezon City read as Quezon province.",
]

# Dropdown options. The value written into the cell begins with the code, so the
# merge step can split on the first " - ".
Q1_OPTS = [
    "Keep - genuine food-insecurity evidence",
    "Drop - not about food insecurity",
    "Unsure - send to adjudicator",
]
Q2_OPTS = [f"{k} - {v[0]}" for k, v in DIMENSIONS.items()] + ["Cannot tell"]
Q3_OPTS = [f"{k} - {HYPOTHESIS_GLOSS[k]}" for k in HYPOTHESIS_TEXT] + [
    "None of these", "Cannot tell"]
Q4_OPTS = [s[0] for s in SEVERITY]
Q5_OPTS = ["Yes - location is right", "No - wrong location", "Cannot tell"]
FLAG_OPTS = ["No", "Yes - needs adjudication"]

QUESTIONS = [
    ("q1", "Q1 - Is this about food insecurity?", 30, Q1_OPTS),
    ("q2", "Q2 - Dimension (A-F)", 26, Q2_OPTS),
    ("q3", "Q3 - Hypothesis (T1-T9)", 28, Q3_OPTS),
    ("q4", "Q4 - Severity", 17, Q4_OPTS),
    ("q5", "Q5 - Is the location right?", 21, Q5_OPTS),
    ("flag", "Flag for adjudication?", 15, FLAG_OPTS),
    ("notes", "Notes / rationale", 40, None),
]

# Headline and lead stay frozen while the reviewer answers, so these two widths
# also drive the row-height estimate below.
W_TITLE, W_LEAD = 46, 52

REF_COLS = [
    ("#", 5, "_n"),
    ("Headline", W_TITLE, "title"),
    ("Opening lines (lead)", W_LEAD, "_lead"),
    ("Date", 11, "_date"),
    ("Date basis", 11, "_date_basis"),
    ("Province", 13, "province"),
    ("City / Municipality", 18, "city_municipality"),
    ("Scope", 15, "_scope"),
    ("Publisher", 17, "news_source"),
    ("Link", 14, "_link"),
]
SYS_COLS = [
    ("SYS: Dimension", 13, "food_security_dimension"),
    ("SYS: Hypothesis", 14, "hypothesis_topic"),
    ("SYS: Category", 26, "food_insecurity_category"),
    ("SYS: Tier", 11, "relevance_tier"),
    ("SYS: Why the audit kept it", 32, "review_basis"),
]

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
NAVY = "1F3A5F"
GOLD = "B8860B"
HEAD_REF = PatternFill("solid", fgColor=NAVY)
HEAD_SYS = PatternFill("solid", fgColor="6E7781")
HEAD_Q = PatternFill("solid", fgColor=GOLD)
FILL_SYS = PatternFill("solid", fgColor="EDEDED")
FILL_Q = PatternFill("solid", fgColor="FFF8E1")
FILL_BAND = PatternFill("solid", fgColor="F6F8FB")
FILL_INPUT = PatternFill("solid", fgColor="FFF3C4")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
UNLOCKED = Protection(locked=False)

DATE_BASIS = {"iso": "exact", "rfc": "feed date", "recovered": "recovered"}
SCOPE_LABEL = {
    "city_municipality": "city / municipality",
    "province": "province-level",
    "region": "region-wide",
    "national": "nationwide",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def clean(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def sentence(text: str) -> str:
    """Capitalise the first letter so a table cell does not open lower-case."""
    return text[:1].upper() + text[1:] if text else text


def est_height(text: str, chars_per_line: int, base: int = 15) -> float:
    cpl = max(1, chars_per_line)
    lines = sum(max(1, math.ceil(len(seg) / cpl)) for seg in str(text).split("\n"))
    return min(base * max(1, lines), 300)


def lead_text(row) -> tuple[str, bool]:
    """The lead, or an honest placeholder. Second value flags the placeholder."""
    lead = clean(row.get("content_lead"))
    if lead:
        return lead, False
    if clean(row.get("title_provenance")) == "slug_derived":
        return ("(No lead text available, and this headline was rebuilt from the article "
                "link, so its wording is normalised rather than verbatim. Judge from the "
                "headline, the date and the location.)", True)
    return ("(No lead text available. Judge from the headline, the date and the location.)",
            True)


def validate_label_space(df: pd.DataFrame) -> None:
    bad_h = sorted(set(df["hypothesis_topic"].dropna()) - set(HYPOTHESIS_TEXT))
    bad_d = sorted(set(df["food_security_dimension"].dropna()) - set(DIMENSIONS))
    if bad_h or bad_d:
        raise SystemExit(
            f"Label-space drift: unknown hypotheses {bad_h}, unknown dimensions {bad_d}. "
            "Update the maps at the top of this script to match audit_stage2_rebuild.py."
        )


def stratified_sample(df: pd.DataFrame, by: str, n: int, seed: int) -> pd.DataFrame:
    """A proportional sample of n rows, trimmed or topped up to land exactly on n."""
    if n <= 0 or df.empty:
        return df.iloc[0:0]
    df = df.reset_index(drop=True)
    n = min(n, len(df))
    s = df.groupby(by, group_keys=False, dropna=False, observed=True).sample(
        frac=n / len(df), random_state=seed)
    if len(s) < n:
        rest = df.drop(index=s.index)
        s = pd.concat([s, rest.sample(n - len(s), random_state=seed)])
    elif len(s) > n:
        s = s.sample(n, random_state=seed)
    return s.reset_index(drop=True)


# ---------------------------------------------------------------------------
# Prose sheets
# ---------------------------------------------------------------------------
def write_prose(ws, lines: list[tuple], b_width: int = 26, c_width: int = 96) -> None:
    """lines: (kind, text) for full-width rows, or (kind, label, text) for two-column."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = b_width
    ws.column_dimensions["C"].width = c_width

    r = 1
    for item in lines:
        kind = item[0]
        if kind == "sp":
            ws.row_dimensions[r].height = 8
            r += 1
            continue

        if len(item) == 3:  # two-column row
            label, text = item[1], item[2]
            cl = ws.cell(row=r, column=2, value=label)
            ct = ws.cell(row=r, column=3, value=text)
            cl.alignment = WRAP
            ct.alignment = WRAP
            if kind == "thead":
                for c in (cl, ct):
                    c.font = HEAD_FONT
                    c.fill = HEAD_REF
                    c.border = BORDER
                ws.row_dimensions[r].height = 18
            else:
                cl.font = Font(bold=True, size=10, color=NAVY)
                ct.font = Font(size=10, color="333333")
                for c in (cl, ct):
                    c.border = BORDER
                ws.row_dimensions[r].height = est_height(text, c_width - 4)
            r += 1
            continue

        text = item[1]
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        cell = ws.cell(row=r, column=2, value=text)
        cell.alignment = WRAP
        if kind == "h1":
            cell.font = Font(bold=True, size=16, color=NAVY)
            ws.row_dimensions[r].height = 22
        elif kind == "h2":
            cell.font = Font(bold=True, size=12, color=GOLD)
            ws.row_dimensions[r].height = 18
        elif kind == "h3":
            cell.font = Font(bold=True, size=11, color=NAVY)
            ws.row_dimensions[r].height = 20
        elif kind == "b":
            cell.font = Font(bold=True, size=10, color="222222")
            ws.row_dimensions[r].height = est_height(text, b_width + c_width - 6)
        else:  # p
            cell.font = Font(size=10, color="333333")
            ws.row_dimensions[r].height = est_height(text, b_width + c_width - 6)
        r += 1


def sheet_start_here(wb, reviewer: str, n_blind: int, n_review: int) -> None:
    ws = wb.create_sheet("Start here")
    write_prose(ws, [
        ("h1", "Expert Annotation Workbook"),
        ("h2", f"aiPHeed CALABARZON food-insecurity news dataset  ·  Reviewer {reviewer}"),
        ("sp",),
        ("p", "Thank you for reviewing this dataset. A machine pipeline collected these "
              "news articles and tagged each one; your job is to judge them independently. "
              "Your answers become the gold standard the system's accuracy is measured "
              "against, and the agreement between reviewers is what makes that measurement "
              "credible. Disagreeing with the machine is not a problem — it is the point."),
        ("sp",),
        ("h3", "What to do, in order"),
        ("b", "1.  Fill in the 'Annotator' tab first."),
        ("p", "     Your credentials and an independence declaration. This is reported in "
              "the study's validity section, so please complete every field."),
        ("b", "2.  Read the 'Codebook' tab once before you start."),
        ("p", "     It defines every option in the drop-downs, including the boundary calls "
              "that cause most disagreement. Ten minutes here saves an hour later."),
        ("b", f"3.  Do 'Step 1 - Blind set' next ({n_blind} articles). Do this one FIRST."),
        ("p", "     This tab deliberately shows you no machine labels, so your judgment is "
              "not anchored to what the pipeline guessed. It is a mixed sample drawn from "
              "both the articles our audit kept and the articles our audit removed, and "
              "they are not marked — so please judge each one on its own merits. This is "
              "how we find articles we wrongly threw away, not just ones we wrongly kept."),
        ("b", f"4.  Then do 'Step 2 - Review set' ({n_review} articles)."),
        ("p", "     Same questions. Here the grey SYS columns show what the pipeline "
              "proposed, so you are checking pre-filled work rather than starting cold. "
              "Your answer overrides the machine's whenever they differ."),
        ("sp",),
        ("h3", "How to fill it in"),
        ("p", "Fill only the gold-headed columns on the right. Each one has a drop-down: "
              "click the cell, then pick a value. The 'Progress' tab counts how far you "
              "have got."),
        ("p", "If Q1 is 'Drop', leave Q2, Q3 and Q4 blank — there is no dimension or "
              "severity to record for an article that does not belong in the dataset. "
              "Still answer Q5 if you can see the location is wrong."),
        ("sp",),
        ("h3", "Ground rules"),
        ("p", "•  Judge the article as written. Do not fill gaps from your own knowledge of "
              "the event, even where you know more than the reporter did."),
        ("p", "•  One primary dimension per article. Many articles touch two; pick the one "
              "the article is most about, and say so in the notes if it was close."),
        ("p", "•  When you are unsure, set 'Flag for adjudication' to Yes rather than "
              "guessing. Flagged rows go to the adjudicator, which is a better outcome for "
              "the dataset than a confident wrong label."),
        ("p", "•  Please do not discuss individual articles with the other reviewer until "
              "both workbooks are complete. Independent answers are what let us compute an "
              "agreement statistic; compared notes destroy it."),
        ("p", "•  Many links point at Google News and no longer open — roughly three in "
              "four articles here have no lead text saved for that reason. That is a known "
              "limit of the collection, not a task for you to fix. Judge from the headline, "
              "date and location, and use 'Cannot tell' when the headline is genuinely too "
              "thin."),
        ("p", "•  Notes are optional but valuable, especially on anything you flag or where "
              "you overrule the machine. One line is enough."),
        ("sp",),
        ("h3", "A note on the locked cells"),
        ("p", "The article columns are locked so they cannot be edited by accident — the "
              "text has to stay identical across reviewers for the answers to be "
              "comparable. There is no password: if you genuinely need to change something, "
              "use Review > Unprotect Sheet in Excel."),
        ("sp",),
        ("p", "Codebook and definitions: the 'Codebook' tab.  ·  Questions: the aiPHeed "
              "research team."),
    ])


def sheet_annotator(wb, reviewer: str) -> None:
    ws = wb.create_sheet("Annotator")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 62

    rows = [
        ("h1", "Annotator record"),
        ("p", "Reported in the study's validity and reliability section. Please complete "
              "every field — an examiner is more likely to ask about this page than about "
              "any single article."),
        ("sp",),
    ]
    write_prose(ws, rows, b_width=42, c_width=62)

    fields = [
        ("Reviewer code", reviewer, False),
        ("Full name", "", True),
        ("Position / title", "", True),
        ("Institution or office", "", True),
        ("Field of expertise", "", True),
        ("Highest relevant degree", "", True),
        ("Years of relevant experience", "", True),
        ("Email address", "", True),
        ("Date started (yyyy-mm-dd)", "", True),
        ("Date completed (yyyy-mm-dd)", "", True),
        ("Did you take any part in building aiPHeed's classifier, or in choosing its "
         "thresholds or labels?", "", True),
        ("Any other interest to declare that a reader should know about?", "", True),
        ("Signature (type your full name)", "", True),
    ]
    r = 5
    for label, value, editable in fields:
        cl = ws.cell(row=r, column=2, value=label)
        cl.font = Font(bold=True, size=10, color=NAVY)
        cl.alignment = WRAP
        cl.border = BORDER
        cv = ws.cell(row=r, column=3, value=value)
        cv.alignment = WRAP
        cv.border = BORDER
        cv.font = Font(size=10)
        if editable:
            cv.fill = FILL_INPUT
        ws.row_dimensions[r].height = est_height(label, 40, base=16)
        r += 1

    dv = DataValidation(type="list", formula1='"No,Yes"', allow_blank=True,
                        showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"C{5 + 10}")  # the independence question


def sheet_codebook(wb) -> None:
    ws = wb.create_sheet("Codebook")
    lines: list[tuple] = [
        ("h1", "Codebook"),
        ("p", "Definitions for every drop-down option. The unit of judgment is one news "
              "article, judged as written."),
        ("sp",),

        ("h3", "Q1  Is this about food insecurity?"),
        ("p", "One inclusion rule: the article must report a food-system condition, cause, "
              "or consequence. A hazard on its own does not qualify."),
        ("thead", "Answer", "When to use it"),
        ("row", "Keep", "The article reports evidence about food access, supply, hunger, "
                        "nutrition, farming, fisheries or food aid — or about a "
                        "determinant of one of those, such as lost income or a disaster "
                        "that cut off supply."),
        ("row", "Drop", "The food connection is incidental or absent. A restaurant or "
                        "recipe feature, a 'solar farm', an agri-tourism or festival "
                        "piece, a research announcement, corporate CSR or celebrity "
                        "relief publicity. Also drop disaster reporting with no "
                        "food-system content: casualties, debris, rainfall totals, class "
                        "suspensions, or a bare state-of-calamity declaration."),
        ("row", "Unsure", "You can argue it either way. Say why in the notes."),
        ("sp",),
        ("b", "The line that decides most typhoon stories"),
        ("p", "Keep it when the article reports crop or fishery damage, evacuation or "
              "displacement, or relief distribution. Drop it when the article reports only "
              "the hazard itself — how much rain fell, how many died, which classes were "
              "suspended, that a calamity was declared."),
        ("sp",),

        ("h3", "Q2  Dimension"),
        ("p", "Pick the ONE dimension the article is most about. The question in the middle "
              "column is the test to apply."),
        ("thead", "Dimension", "Ask yourself"),
    ]
    for code, (name, question, detail) in DIMENSIONS.items():
        lines.append(("row", f"{code}  ·  {name}", f"{question}\n{detail}"))
    lines += [
        ("row", "Cannot tell", "The article is too thin to place. Prefer this over a guess."),
        ("sp",),
        ("b", "Boundary calls — where reviewers most often disagree"),
        ("thead", "Pair", "What separates them"),
    ]
    for pair, text in BOUNDARIES:
        lines.append(("row", pair, sentence(text)))
    lines += [
        ("sp",),
        ("h3", "Q3  Hypothesis"),
        ("p", "The thesis tests ten hypotheses. Pick the one the article gives evidence "
              "for. The dimension each hypothesis belongs to is shown so you can check "
              "your Q2 answer against it — if your Q2 and Q3 point at different "
              "dimensions, one of them needs another look."),
        ("thead", "Code", "Hypothesis"),
    ]
    for code, text in HYPOTHESIS_TEXT.items():
        lines.append(("row", f"{code}  (dimension {HYPOTHESIS_DIM[code]})", text))
    lines += [
        ("row", "None of these", "The article is about food insecurity but none of the ten "
                                 "fits. Please say what it is about in the notes — these "
                                 "are the most useful rows for improving the scheme."),
        ("row", "Cannot tell", "Too thin to place."),
        ("sp",),

        ("h3", "Q4  Severity"),
        ("p", "How severe is the food-insecurity condition this article reports, as the "
              "article reports it? Judge what is on the page, not the wider event."),
        ("thead", "Level", "Meaning"),
    ]
    for level, text in SEVERITY:
        lines.append(("row", level, text))
    lines += [
        ("sp",),
        ("h3", "Q5  Is the location right?"),
        ("p", "The province and city shown were assigned by the pipeline. Earlier versions "
              "of this dataset were assigned lexically and got this wrong in systematic "
              "ways; the audit corrected the cases it caught, and this question is how we "
              "check whether any survived. Answer from the article's subject, not its "
              "dateline."),
        ("thead", "Look for", "Example"),
    ]
    for trap in LOCATION_TRAPS:
        head, _, rest = trap.partition(" — ")
        lines.append(("row", head, sentence(rest)))
    lines += [
        ("sp",),
        ("h3", "Reading the reference columns"),
        ("thead", "Column", "What it means"),
        ("row", "Opening lines (lead)", "The first lines of the article. Empty for about "
                                        "three in four rows — see the note in the cell."),
        ("row", "Date basis", "'exact' = the date came from the article itself. 'feed date' "
                              "= it came from a news feed and may be the crawl date rather "
                              "than the publication date. 'recovered' = it was reconstructed "
                              "from a truncated stamp. Treat the last two as approximate."),
        ("row", "Scope", "Whether the article is about one city or municipality, a whole "
                         "province, or the region. Region-wide rows have no province, which "
                         "is correct rather than missing."),
        ("row", "Link", "Opens the article. Many point at Google News and no longer "
                        "resolve; that is a known collection limit."),
        ("row", "SYS: ... (grey)", "What the pipeline proposed. Present on the review set "
                                   "only, and never on the blind set. Treat it as a "
                                   "suggestion to check, not an answer to confirm."),
    ]
    write_prose(ws, lines)


# ---------------------------------------------------------------------------
# Data sheets
# ---------------------------------------------------------------------------
def build_lists_sheet(wb) -> dict[str, str]:
    ws = wb.create_sheet("Lists")
    refs = {}
    for ci, (key, _, _, opts) in enumerate(
            [q for q in QUESTIONS if q[3] is not None], start=1):
        letter = get_column_letter(ci)
        ws.cell(row=1, column=ci, value=key)
        for ri, opt in enumerate(opts, start=2):
            ws.cell(row=ri, column=ci, value=opt)
        refs[key] = f"Lists!${letter}$2:${letter}${len(opts) + 1}"
    ws.sheet_state = "hidden"
    return refs


def build_data_sheet(wb, title: str, rows: pd.DataFrame, refs: dict[str, str],
                     show_sys: bool) -> dict:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False

    plan = [(h, "ref", w, f) for h, w, f in REF_COLS]
    if show_sys:
        plan += [(h, "sys", w, f) for h, w, f in SYS_COLS]
    plan += [(h, "q", w, key) for key, h, w, _ in QUESTIONS]
    plan += [("Article ID", "ref", 15, "article_id")]

    qcols: dict[str, str] = {}
    for ci, (header, kind, width, field) in enumerate(plan, start=1):
        letter = get_column_letter(ci)
        c = ws.cell(row=1, column=ci, value=header)
        c.font = HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.fill = {"ref": HEAD_REF, "sys": HEAD_SYS, "q": HEAD_Q}[kind]
        c.border = BORDER
        ws.column_dimensions[letter].width = width
        if kind == "q":
            qcols[field] = letter
    ws.row_dimensions[1].height = 34
    # Freeze #, Headline and lead: the reviewer needs the evidence in view while
    # answering the questions off to the right.
    ws.freeze_panes = "D2"

    for ri, (_, row) in enumerate(rows.iterrows(), start=2):
        lead, is_placeholder = lead_text(row)
        band = FILL_BAND if ri % 2 == 0 else None
        for ci, (header, kind, width, field) in enumerate(plan, start=1):
            value = ""
            if kind == "q":
                value = ""
            elif field == "_n":
                value = ri - 1
            elif field == "_lead":
                value = lead
            elif field == "_date":
                value = clean(row.get("publication_date"))[:10]
            elif field == "_date_basis":
                value = DATE_BASIS.get(clean(row.get("date_provenance")), "")
            elif field == "_scope":
                value = SCOPE_LABEL.get(clean(row.get("geographic_scope")), "")
            elif field == "_link":
                url = clean(row.get("url"))
                value = ("Open (Google)" if "news.google.com" in url
                         else "Open article") if url else ""
            else:
                value = clean(row.get(field))

            c = ws.cell(row=ri, column=ci, value=value)
            c.border = BORDER
            c.alignment = WRAP_CENTER if kind == "sys" or field in (
                "_n", "_date", "_date_basis") else WRAP

            if kind == "q":
                c.fill = FILL_Q
                c.protection = UNLOCKED
            elif kind == "sys":
                c.fill = FILL_SYS
                c.font = Font(size=9, color="444444")
            else:
                if band:
                    c.fill = band
                c.font = Font(size=10, color="222222")
                if field == "title":
                    c.font = Font(size=10, bold=True, color="111111")
                elif field == "_lead" and is_placeholder:
                    c.font = Font(size=9, italic=True, color="888888")
                elif field == "article_id":
                    c.font = Font(size=8, color="999999")
                elif field == "_link" and value:
                    c.hyperlink = clean(row.get("url"))
                    c.font = Font(size=10, color="0563C1", underline="single")

        body = lead if not is_placeholder else ""
        # Tall enough to show the whole lead for ~90% of the rows that have one.
        # A clipped lead is worse than a tall row: the lead is the evidence.
        ws.row_dimensions[ri].height = max(
            32.0, min(240.0, 15 + est_height(body, W_LEAD - 4)
                      + est_height(clean(row.get("title")), W_TITLE - 4)))

    last = len(rows) + 1
    for key, _, _, opts in QUESTIONS:
        if opts is None:
            continue
        dv = DataValidation(type="list", formula1=refs[key], allow_blank=True,
                            showErrorMessage=True, errorStyle="stop",
                            errorTitle="Pick from the list",
                            error="Please choose one of the drop-down values. If none of "
                                  "them fits, pick the 'Cannot tell' option and explain in "
                                  "the notes column.")
        ws.add_data_validation(dv)
        dv.add(f"{qcols[key]}2:{qcols[key]}{last}")

    ws.auto_filter.ref = f"A1:{get_column_letter(len(plan))}{last}"
    ws.protection = SheetProtection(
        sheet=True, autoFilter=False, sort=False,
        formatCells=False, formatColumns=False, formatRows=False,
        insertRows=True, insertColumns=True, deleteRows=True, deleteColumns=True,
    )
    return {"title": title, "first": 2, "last": last, "qcols": qcols}


def sheet_progress(wb, sheets: list[dict]) -> None:
    ws = wb.create_sheet("Progress")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16

    ws.cell(row=1, column=2, value="Progress").font = Font(bold=True, size=16, color=NAVY)
    ws.cell(row=2, column=2,
            value="Counts update themselves as you fill the sheets in.").font = Font(
        size=10, color="555555")

    r = 4
    for meta in sheets:
        q1 = meta["qcols"]["q1"]
        total = meta["last"] - meta["first"] + 1
        ref = f"'{meta['title']}'!${q1}${meta['first']}:${q1}${meta['last']}"

        h = ws.cell(row=r, column=2, value=meta["title"])
        h.font = Font(bold=True, size=11, color=NAVY)
        r += 1
        entries = [
            ("Articles in this tab", total),
            ("Answered (Q1 filled)", f"=COUNTA({ref})"),
            ("Still to do", f"=({total})-COUNTA({ref})"),
            ("Kept", f'=COUNTIF({ref},"Keep*")'),
            ("Dropped", f'=COUNTIF({ref},"Drop*")'),
            ("Unsure", f'=COUNTIF({ref},"Unsure*")'),
        ]
        flag = meta["qcols"]["flag"]
        fref = f"'{meta['title']}'!${flag}${meta['first']}:${flag}${meta['last']}"
        entries.append(("Flagged for adjudication", f'=COUNTIF({fref},"Yes*")'))

        for label, value in entries:
            cl = ws.cell(row=r, column=2, value=label)
            cl.font = Font(size=10, color="333333")
            cl.border = BORDER
            cv = ws.cell(row=r, column=3, value=value)
            cv.font = Font(size=10, bold=True)
            cv.alignment = Alignment(horizontal="center")
            cv.border = BORDER
            r += 1
        r += 1


# ---------------------------------------------------------------------------
# Assembly
# ---------------------------------------------------------------------------
def load_frames(args) -> tuple[pd.DataFrame, pd.DataFrame]:
    kept = pd.read_parquet(SRC)
    validate_label_space(kept)

    dropped = pd.DataFrame()
    if args.blind_drop > 0 and DROPS.exists() and STAGE1.exists():
        drops = pd.read_csv(DROPS)
        pool = drops[drops["stage"] == "stage2_relevance_review"].copy()
        s1 = pd.read_parquet(STAGE1)
        keep_cols = [c for c in ("article_id", "title", "publication_date",
                                 "date_provenance", "news_source", "url", "content_lead",
                                 "title_provenance", "province", "city_municipality",
                                 "geographic_scope") if c in s1.columns]
        dropped = pool[["article_id", "drop_reason"]].merge(
            s1[keep_cols], on="article_id", how="inner")
    elif args.blind_drop > 0:
        print("  ! drop-sample sources missing; blind set will contain kept rows only")
    return kept, dropped


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("-r", "--reviewers", nargs="+", default=["A", "B"],
                    help="reviewer codes; one workbook each (default: A B)")
    ap.add_argument("--blind-keep", type=int, default=60,
                    help="kept articles in the blind set (default: 60)")
    ap.add_argument("--blind-drop", type=int, default=40,
                    help="audit-removed articles mixed into the blind set (default: 40)")
    ap.add_argument("--pilot", type=int, default=12,
                    help="rows in the extra pilot workbook; 0 to skip (default: 12)")
    ap.add_argument("--seed", type=int, default=20260821)
    ap.add_argument("--outdir", type=Path, default=OUTDIR)
    args = ap.parse_args()

    kept, dropped = load_frames(args)
    args.outdir.mkdir(parents=True, exist_ok=True)

    # --- blind set: kept sample + removed sample, interleaved, labels withheld
    blind_keep = stratified_sample(kept, "food_security_dimension",
                                   args.blind_keep, args.seed)
    blind_keep["_audit_decision"] = "keep"
    if len(dropped):
        blind_drop = stratified_sample(dropped, "province", args.blind_drop, args.seed)
        blind_drop["_audit_decision"] = "drop"
    else:
        blind_drop = dropped.assign(_audit_decision="drop")
    blind = pd.concat([blind_keep, blind_drop], ignore_index=True)
    blind = blind.sample(frac=1, random_state=args.seed).reset_index(drop=True)

    # --- review set: every kept article the blind set did not take
    review = kept[~kept["article_id"].isin(blind_keep["article_id"])].sort_values(
        ["province", "publication_date"], na_position="last").reset_index(drop=True)

    refs_blind = None
    for reviewer in args.reviewers:
        wb = Workbook()
        wb.remove(wb.active)
        sheet_start_here(wb, reviewer, len(blind), len(review))
        sheet_annotator(wb, reviewer)
        sheet_codebook(wb)
        refs = build_lists_sheet(wb)
        refs_blind = refs
        m1 = build_data_sheet(wb, "Step 1 - Blind set", blind, refs, show_sys=False)
        m2 = build_data_sheet(wb, "Step 2 - Review set", review, refs, show_sys=True)
        sheet_progress(wb, [m1, m2])
        out = args.outdir / f"annotation_reviewer_{reviewer}.xlsx"
        wb.save(out)
        print(f"  wrote {out.relative_to(ROOT)}  "
              f"({len(blind)} blind + {len(review)} review)")

    # --- pilot workbook: a small batch to shake out the codebook before scaling
    if args.pilot > 0:
        pilot = stratified_sample(kept, "food_security_dimension", args.pilot, args.seed + 1)
        wb = Workbook()
        wb.remove(wb.active)
        sheet_start_here(wb, "PILOT", 0, len(pilot))
        sheet_annotator(wb, "PILOT")
        sheet_codebook(wb)
        refs = build_lists_sheet(wb)
        mp = build_data_sheet(wb, "Step 2 - Review set", pilot, refs, show_sys=True)
        sheet_progress(wb, [mp])
        out = args.outdir / "annotation_pilot.xlsx"
        wb.save(out)
        print(f"  wrote {out.relative_to(ROOT)}  ({len(pilot)} articles)")

    # --- answer key for the blind set. Never ships inside a reviewer workbook.
    key_cols = ["article_id", "_audit_decision", "title", "publication_date", "province",
                "city_municipality", "news_source", "url"]
    for c in ("food_security_dimension", "hypothesis_topic", "relevance_tier",
              "review_basis", "drop_reason"):
        if c in blind.columns:
            key_cols.append(c)
    key = blind[[c for c in key_cols if c in blind.columns]].copy()
    key.insert(0, "blind_row", range(1, len(key) + 1))
    key_path = args.outdir / "blind_set_key.csv"
    key.to_csv(key_path, index=False, encoding="utf-8-sig")

    manifest = {
        "generated_from": str(SRC.relative_to(ROOT)),
        "seed": args.seed,
        "reviewers": args.reviewers,
        "kept_articles_total": int(len(kept)),
        "blind_set": {"kept": int(len(blind_keep)), "audit_removed": int(len(blind_drop)),
                      "total": int(len(blind))},
        "review_set": int(len(review)),
        "pilot": int(args.pilot),
        "coverage_note": ("every kept article appears exactly once per reviewer, in "
                          "either the blind set or the review set"),
    }
    (args.outdir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n",
                                               encoding="utf-8")
    print(f"  wrote {key_path.relative_to(ROOT)} (answer key - keep away from reviewers)")
    print(f"  wrote {(args.outdir / 'manifest.json').relative_to(ROOT)}")


if __name__ == "__main__":
    main()
