"""
build_relevance_sample.py
--------------------------
Draw and score a labelling sample for evaluating the news relevance gate.

The problem this answers: the zero-shot NLI classifier sometimes collapses into
entailing every hypothesis at once, so an ATM press release scores 0.9943 on
"fish kills" -- the argmax of noise rather than a detection. No absolute
threshold separates that from a real fish-kill story scoring 0.21. Choosing a
better gate needs measured precision and recall, not a heuristic picked by
eye.

Two arms, because a precision-only sample cannot see what the gate is throwing
away:

  KEPT      stratified across the score range of the scored corpus. Answers
            "of what we publish, how much is junk?"
  REJECTED  drawn from the pre-classification pool, excluding anything that
            reached the corpus. Answers "of what we discard, how much was
            real?" Without this arm every gate looks better the stricter it
            gets, because tightening can only raise precision.

Both arms are re-scored against all ten hypotheses and the vectors are cached,
so evaluate_relevance_gates.py can compare candidate gates without touching
the model again.

The workbook is BLIND: it shows the labeller the headline and lead only. If it
showed the machine's score or its predicted topic, the labels would anchor to
them and the measurement would flatter the classifier. The arm and the scores
live in the key file instead.

Usage
-----
    python scripts/build_relevance_sample.py                    # 150 kept, 100 rejected
    python scripts/build_relevance_sample.py --kept 200 --rejected 150
    python scripts/build_relevance_sample.py --resume           # keep existing scores
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from app.ml.nlp.classifier import (  # noqa: E402
    CORE_TOPIC_IDS,
    HYPOTHESES,
    RELEVANCE_THRESHOLD,
    load_classifier,
)

SCORED = ROOT / "data" / "processed" / "corpus_geocoded.parquet"
POOL = ROOT / "data" / "processed" / "_fullpool_geocoded.parquet"
OUTDIR = ROOT / "data" / "processed" / "relevance_eval"

FEATURES_PATH = OUTDIR / "sample_features.parquet"
WORKBOOK_PATH = OUTDIR / "relevance_labelling.xlsx"

# Bins spanning the scored range. Equal draws per bin oversample the sparse
# high end on purpose: that is where the gate's decision boundary sits and
# where a handful of labels move precision the most.
SCORE_BINS = [0.30, 0.50, 0.70, 0.90, 0.95, 1.01]

RELEVANCE_OPTIONS = [
    "Yes - about food access, supply, prices, hunger or farm/fishery output",
    "No - not about food security",
    "Unsure",
]

TOPIC_OPTIONS = ["(not relevant)"] + [
    f"{tid} - {text.replace('This article is about ', '')}"
    for tid, text in HYPOTHESES.items()
] + ["Other food topic not listed"]

HEADER_FILL = PatternFill("solid", fgColor="1F3A2C")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF7E0")


# ---------------------------------------------------------------------------
# Sampling
# ---------------------------------------------------------------------------

def draw_kept(n: int, seed: int) -> pd.DataFrame:
    """Stratified draw across score bins from the published corpus."""
    df = pd.read_parquet(SCORED)
    df = df[df["food_insecurity_score"].notna()].copy()
    df["stratum"] = pd.cut(df["food_insecurity_score"], SCORE_BINS, right=False)

    per_bin = max(n // (len(SCORE_BINS) - 1), 1)
    parts = []
    for stratum, group in df.groupby("stratum", observed=True):
        parts.append(group.sample(min(per_bin, len(group)), random_state=seed))
    out = pd.concat(parts)
    # Interval objects do not survive a parquet round-trip; the label is only
    # ever read back as text anyway.
    out["stratum"] = out["stratum"].astype(str)
    out["arm"] = "kept"
    return out[["article_id", "title", "summary", "link", "published",
                "source_domain", "food_insecurity_score", "top_hypothesis",
                "stratum", "arm"]]


def draw_rejected(n: int, seed: int) -> pd.DataFrame:
    """
    Random draw from the pre-classification pool, excluding the published set.

    Unstratified on purpose: there is no score to stratify on, and the point of
    this arm is an unbiased estimate of how much genuine signal the gate is
    discarding.
    """
    pool = pd.read_parquet(POOL)
    scored_ids = set(pd.read_parquet(SCORED)["article_id"])
    pool = pool[~pool["article_id"].isin(scored_ids)]
    pool = pool[pool["title"].notna() & (pool["title"].str.len() > 15)]

    out = pool.sample(min(n, len(pool)), random_state=seed).copy()
    out["link"] = None
    out["published"] = None
    out["source_domain"] = None
    out["food_insecurity_score"] = None
    out["top_hypothesis"] = None
    out["stratum"] = "rejected-pool"
    out["arm"] = "rejected"
    return out[["article_id", "title", "summary", "link", "published",
                "source_domain", "food_insecurity_score", "top_hypothesis",
                "stratum", "arm"]]


# ---------------------------------------------------------------------------
# Scoring
# ---------------------------------------------------------------------------

def score_all(df: pd.DataFrame) -> pd.DataFrame:
    """
    Attach the full ten-hypothesis entailment vector to every row.

    Cached so gate evaluation never needs the model: on CPU this is the slow
    step, roughly a second per article.
    """
    clf = load_classifier()
    if getattr(clf, "mode", None) != "xlm-roberta":
        raise SystemExit(
            "The NLI model did not load, so the sample would be scored by the "
            "keyword fallback and measure the wrong thing. Fix the model load "
            "before drawing a sample."
        )

    ids = list(HYPOTHESES.keys())
    rows = []
    total = len(df)
    for i, (_, r) in enumerate(df.iterrows(), 1):
        text = f"{r['title']}. {r['summary'] or ''}"[:512]
        scores = {d["topic_id"]: round(d["score"], 6) for d in clf.classify(text)}
        core = {k: v for k, v in scores.items() if k in CORE_TOPIC_IDS}
        noncore = {k: v for k, v in scores.items() if k not in CORE_TOPIC_IDS}
        rows.append({
            **{f"h_{tid}": scores[tid] for tid in ids},
            "core_max": max(core.values()),
            "core_top": max(core, key=core.get),
            "noncore_max": max(noncore.values()),
            "n_over_090": sum(1 for v in scores.values() if v >= 0.90),
        })
        if i % 25 == 0 or i == total:
            print(f"  scored {i}/{total}", flush=True)

    return pd.concat([df.reset_index(drop=True), pd.DataFrame(rows)], axis=1)


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def _write_instructions(wb: Workbook, n_rows: int) -> None:
    ws = wb.create_sheet("Start here", 0)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 100

    lines = [
        ("h", "Labelling the news relevance sample"),
        ("", ""),
        ("p", f"There are {n_rows} articles on the 'Label' sheet. For each one, "
              "decide whether it carries information about food security in "
              "CALABARZON, then pick the closest topic."),
        ("", ""),
        ("h2", "What counts as relevant"),
        ("p", "Yes - the article is about food prices or supply, hunger or "
              "malnutrition, food assistance, crop or fishery production, "
              "farmland loss, or an event clearly disrupting food access "
              "(a typhoon damaging harvests, a fish kill, a transport failure "
              "affecting food delivery)."),
        ("p", "No - the article is about something else. Banking, sport, "
              "showbiz, general politics, crime, infrastructure with no food "
              "angle. A passing mention of the word 'food' is not enough."),
        ("p", "Unsure - genuinely ambiguous. Use it sparingly; it is excluded "
              "from precision and recall rather than counted either way."),
        ("", ""),
        ("h2", "Two things to keep in mind"),
        ("p", "Judge the article, not the region. Some articles are national "
              "or from other provinces; label relevance to food security, and "
              "let the geocoder worry about location."),
        ("p", "You are not being shown the machine's score or its guess. That "
              "is deliberate - if you could see them your labels would drift "
              "toward agreeing with the classifier, and the measurement would "
              "flatter it. The sheet also mixes articles the system kept with "
              "articles it discarded, in no particular order."),
        ("", ""),
        ("h2", "When you are done"),
        ("p", "Save the file in place and run:"),
        ("c", "python scripts/evaluate_relevance_gates.py"),
        ("p", "That reports precision and recall for each candidate gate and "
              "estimates how many articles each would keep corpus-wide."),
    ]

    row = 2
    for kind, text in lines:
        cell = ws.cell(row=row, column=2, value=text)
        if kind == "h":
            cell.font = Font(bold=True, size=15)
        elif kind == "h2":
            cell.font = Font(bold=True, size=11)
        elif kind == "c":
            cell.font = Font(name="Consolas", size=10)
        else:
            cell.font = Font(size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = max(15, 15 * (len(text) // 95 + 1))
        row += 1


def _write_labels(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Label")
    # Source and date are deliberately absent. The pre-classification pool
    # carries neither, so showing them would leave those columns blank on
    # exactly the rejected-arm rows -- telling the labeller which arm each row
    # came from and destroying the blinding. Neither is needed to judge whether
    # an article is about food.
    headers = ["#", "Headline", "Lead", "Relevant?", "Topic", "Notes"]
    widths = [5, 62, 84, 34, 40, 34]

    for col, (head, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=head)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    for i, (_, r) in enumerate(df.iterrows(), start=2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=str(r["title"] or "")).alignment = \
            Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=3, value=str(r["summary"] or "")[:600]).alignment = \
            Alignment(wrap_text=True, vertical="top")
        for col in (4, 5, 6):
            ws.cell(row=i, column=col).fill = INPUT_FILL
        ws.row_dimensions[i].height = 46

    last = len(df) + 1
    dv_rel = DataValidation(
        type="list", formula1='"' + ",".join(RELEVANCE_OPTIONS) + '"',
        allow_blank=True, showDropDown=False)
    dv_rel.error = "Pick one of the listed options."
    ws.add_data_validation(dv_rel)
    dv_rel.add(f"D2:D{last}")

    # Excel caps an inline list at 255 characters, so the topic options live on
    # their own sheet and the validation points at the range.
    lists = wb.create_sheet("_lists")
    for j, opt in enumerate(TOPIC_OPTIONS, start=1):
        lists.cell(row=j, column=1, value=opt)
    lists.sheet_state = "hidden"

    dv_topic = DataValidation(
        type="list", formula1=f"_lists!$A$1:$A${len(TOPIC_OPTIONS)}",
        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_topic)
    dv_topic.add(f"E2:E{last}")


def build_workbook(df: pd.DataFrame, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    _write_instructions(wb, len(df))
    _write_labels(wb, df)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("Usage")[0].strip())
    ap.add_argument("--kept", type=int, default=150,
                    help="articles drawn from the published corpus")
    ap.add_argument("--rejected", type=int, default=100,
                    help="articles drawn from the pre-classification pool")
    ap.add_argument("--seed", type=int, default=20260902)
    ap.add_argument("--resume", action="store_true",
                    help="reuse existing scored sample; only rebuild the workbook")
    args = ap.parse_args()

    OUTDIR.mkdir(parents=True, exist_ok=True)

    if args.resume and FEATURES_PATH.exists():
        scored = pd.read_parquet(FEATURES_PATH)
        print(f"resumed {len(scored)} scored rows from {FEATURES_PATH.name}")
    else:
        sample = pd.concat([
            draw_kept(args.kept, args.seed),
            draw_rejected(args.rejected, args.seed),
        ], ignore_index=True)
        print(f"drawn: {len(sample)} articles "
              f"({(sample['arm'] == 'kept').sum()} kept, "
              f"{(sample['arm'] == 'rejected').sum()} rejected-pool)")
        print("scoring against all 10 hypotheses (slow on CPU)...")
        scored = score_all(sample)
        scored.to_parquet(FEATURES_PATH, index=False)
        print(f"wrote {FEATURES_PATH}")

    # Shuffle so the two arms interleave and the labeller cannot infer which
    # is which from position. On --resume the row numbers are already assigned;
    # renumbering would silently detach existing labels from their articles.
    if "row_no" in scored.columns:
        shuffled = scored.sort_values("row_no").reset_index(drop=True)
    else:
        shuffled = scored.sample(frac=1, random_state=args.seed).reset_index(drop=True)
        shuffled.insert(0, "row_no", range(1, len(shuffled) + 1))
    shuffled.to_parquet(FEATURES_PATH, index=False)

    build_workbook(shuffled, WORKBOOK_PATH)
    print(f"wrote {WORKBOOK_PATH}")
    print()
    print(f"Current gate keeps core_max >= {RELEVANCE_THRESHOLD}. "
          "Label the workbook, then run scripts/evaluate_relevance_gates.py.")


if __name__ == "__main__":
    main()
